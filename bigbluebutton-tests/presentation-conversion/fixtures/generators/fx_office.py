"""Office fixtures (off-*). bbb-web converts them to PDF with LibreOffice first.

Every package is rewritten by pcf_common: fixed core properties (2020-01-01T00:00:00Z,
lastModifiedBy pcf), entries sorted ([Content_Types].xml first) or, for ODF, `mimetype`
first and stored, 1980-01-01 zip timestamps and deflate level 9.
"""
import io

from docx import Document
from docx.shared import Mm, Pt as DocxPt
from odf import dc, meta
from odf.draw import Ellipse, Frame, Line, Page as DrawPage, Rect, TextBox
from odf.opendocument import (OpenDocumentDrawing, OpenDocumentPresentation,
                              OpenDocumentSpreadsheet, OpenDocumentText)
from odf.style import (GraphicProperties, MasterPage, PageLayout, PageLayoutProperties,
                       ParagraphProperties, Style, TableProperties, TextProperties)
from odf.table import Table, TableCell, TableColumn, TableRow
from odf.text import H, P
from openpyxl import Workbook
from openpyxl.styles import Font
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from fx_pdf import camera_exif6_jpeg
from pcf_common import DOC_AUTHOR, normalize_ooxml, write_zip, zip_entries
from pcf_pdfkit import BLUE, GREEN, RED, YELLOW

ODF_DATE = "2020-01-01T00:00:00"
LOREM_PARAS = [
    "Lorem ipsum dolor sit amet, consectetur adipiscing elit. Integer posuere erat a ante "
    "venenatis dapibus posuere velit aliquet. Cras mattis consectetur purus sit amet fermentum. "
    "Donec ullamcorper nulla non metus auctor fringilla. Vestibulum id ligula porta felis euismod "
    "semper. Maecenas faucibus mollis interdum.",
    "Aenean lacinia bibendum nulla sed consectetur. Nullam quis risus eget urna mollis ornare vel "
    "eu leo. Etiam porta sem malesuada magna mollis euismod. Sed posuere consectetur est at "
    "lobortis. Curabitur blandit tempus porttitor. Morbi leo risus, porta ac consectetur ac, "
    "vestibulum at eros.",
    "Praesent commodo cursus magna, vel scelerisque nisl consectetur et. Duis mollis, est non "
    "commodo luctus, nisi erat porttitor ligula, eget lacinia odio sem nec elit. Vivamus sagittis "
    "lacus vel augue laoreet rutrum faucibus dolor auctor. Nulla vitae elit libero, a pharetra augue.",
]
TABLE = [
    ["Fruit", "Colour", "Count", "Price"],
    ["Apple", "Red", 12, 0.5],
    ["Banana", "Yellow", 6, 0.25],
    ["Cherry", "Dark red", 40, 0.1],
    ["Grape", "Green", 25, 0.05],
    ["Lemon", "Yellow", 3, 0.4],
]
OFFICE_PRECHECK = ["defaults", "office-precheck"]


def _hex(rgb):
    return "#%02x%02x%02x" % rgb


def _odf_bytes(doc):
    """odfpy package rewritten: `mimetype` first and stored, the rest sorted."""
    doc.meta.addElement(meta.CreationDate(text=ODF_DATE))
    doc.meta.addElement(dc.Date(text=ODF_DATE))
    doc.meta.addElement(meta.InitialCreator(text=DOC_AUTHOR))
    doc.meta.addElement(dc.Creator(text=DOC_AUTHOR))
    buf = io.BytesIO()
    doc.write(buf)
    entries = zip_entries(buf.getvalue())
    return write_zip(entries, ["mimetype"] + sorted(n for n in entries if n != "mimetype"))


# --------------------------------------------------------------------------------------
# 33. RTF
# --------------------------------------------------------------------------------------
RTF = (
    "{\\rtf1\\ansi\\ansicpg1252\\deff0\n"
    "{\\fonttbl{\\f0\\fswiss\\fcharset0 Arial;}{\\f1\\froman\\fcharset0 Times New Roman;}}\n"
    "{\\info{\\title pcf RTF fixture}{\\author pcf}"
    "{\\creatim\\yr2020\\mo1\\dy1\\hr0\\min0}{\\revtim\\yr2020\\mo1\\dy1\\hr0\\min0}}\n"
    "\\paperw11906\\paperh16838\\margl1418\\margr1418\\margt1418\\margb1418\n"
    "\\pard\\sa240\\f0\\fs48\\b RTF heading\\b0\\par\n"
    "\\pard\\sa120\\f1\\fs28 First short line.\\par\n"
    "Second short line.\\par\n"
    "Third short line.\\par\n"
    "}\n"
)


def off_rtf_one_page(ctx):
    fid = "off-rtf-one-page"
    ctx.write(fid + ".rtf", RTF.encode("ascii"))
    ctx.manifest(
        id=fid, title="Hand-written RTF: a heading and three short lines on A4",
        file=fid + ".rtf", issues=[25432], catalog=["Office pre-check", "LibreOffice process failures"],
        pages=1, expect="converted", kind="any", marker=False, aspect=0.7071, profiles=OFFICE_PRECHECK,
        notes="Plain RTF 1 with an A4 paper size (11906 x 16838 twips). office-precheck sets "
              "skipOfficePrecheck=false, whose inverted isValid() branch (#25432) rejects every "
              "non-pptx file. Measured with LibreOffice 25.8 (bbb-soffice image, file renamed to `file` as "
              "convert-local.sh does): one A4 page, TrueType subsets, no <image or <filter.")


# --------------------------------------------------------------------------------------
# 34. PPTX marker slide
# --------------------------------------------------------------------------------------
def _pptx_band(slide, x, y, w, h, rgb):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(x), Emu(y), Emu(w), Emu(h))
    s.fill.solid()
    s.fill.fore_color.rgb = RGBColor(*rgb)
    s.line.fill.background()
    style = s._element.find(qn("p:style"))      # drop theme line/effect references (shadows)
    if style is not None:
        s._element.remove(style)
    return s


def off_pptx_one_slide(ctx):
    fid = "off-pptx-one-slide"
    prs = Presentation()
    sw, sh = 12192000, 6858000                   # 13.333 x 7.5 in, 16:9
    prs.slide_width, prs.slide_height = Emu(sw), Emu(sh)
    slide = prs.slides.add_slide(prs.slide_layouts[5])   # "Title Only"
    bx, by, gap = sw // 7, sh // 7, sh // 80
    _pptx_band(slide, 0, 0, sw, by, RED)
    _pptx_band(slide, 0, sh - by, sw, by, BLUE)
    _pptx_band(slide, 0, by + gap, bx, sh - 2 * by - 2 * gap, GREEN)
    _pptx_band(slide, sw - bx, by + gap, bx, sh - 2 * by - 2 * gap, YELLOW)
    inner_x, inner_w = bx + Inches(0.3), sw - 2 * bx - 2 * Inches(0.3)
    title = slide.shapes.title
    title.left, title.top, title.width, title.height = inner_x, by + Inches(0.4), inner_w, Inches(1.2)
    title.text = "PPTX marker slide"
    title.text_frame.paragraphs[0].runs[0].font.size = Pt(40)
    box = slide.shapes.add_textbox(inner_x, by + Inches(2.0), inner_w, Inches(1.5))
    box.text_frame.word_wrap = True
    box.text_frame.text = "A title, this text box and four coloured rectangles."
    box.text_frame.paragraphs[0].runs[0].font.size = Pt(24)
    buf = io.BytesIO()
    prs.save(buf)
    ctx.write(fid + ".pptx", normalize_ooxml(buf.getvalue(), "PPTX marker slide"))
    prof = ["defaults", "rasterize-all", "office-precheck"]
    ctx.manifest(
        id=fid, title="python-pptx 16:9 slide: title, text box, red/blue/green/yellow edge rectangles",
        file=fid + ".pptx", issues=[25432], catalog=["Office pre-check"],
        pages=1, expect="converted", kind={"defaults": "any", "rasterize-all": "raster", "office-precheck": "any"},
        marker=True, aspect=1.7778, profiles=prof,
        notes="Layout 'Title Only' of python-pptx's default template, slide size 12192000 x 6858000 "
              "EMU. The four rectangles are laid out as the marker (red top, blue bottom, green "
              "left, yellow right, solid fill, no outline, theme style removed so no shadow), with "
              "the title and the text box in the white middle. A valid pptx: with "
              "skipOfficePrecheck=false (office-precheck) it must still convert (#25432). LibreOffice 25.8: one "
              "960x540 page with the marker upright, 43 <path, 68 <use, no <filter.")


# --------------------------------------------------------------------------------------
# 35 and 41. DOCX
# --------------------------------------------------------------------------------------
def _docx_lorem():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Mm(25)
    doc.add_heading("Lorem ipsum", level=1)
    for text in LOREM_PARAS:
        para = doc.add_paragraph(text)
        para.runs[0].font.size = DocxPt(12)
    buf = io.BytesIO()
    doc.save(buf)
    return normalize_ooxml(buf.getvalue(), "Lorem ipsum A4")


def off_docx_lorem_a4(ctx):
    fid = "off-docx-lorem-a4"
    ctx.docx = ctx.write(fid + ".docx", _docx_lorem())
    ctx.manifest(
        id=fid, title="python-docx A4 page: a heading and three paragraphs of lorem ipsum",
        file=fid + ".docx", issues=[21954, 25432], catalog=["docx text rasterized", "Office pre-check"],
        pages=1, expect="converted", kind="vector", marker=False, aspect=0.7071, profiles=OFFICE_PRECHECK,
        notes="Text-only portrait page. LibreOffice embeds TrueType subsets, so the slide must stay "
              "vector: a raster slide here is #21954. office-precheck must still convert it (#25432). "
              "LibreOffice 25.8: one A4 page, Carlito and Caladea subsets, 42 <path, 887 <use, no <image.")


def off_docx_rezipped(ctx):
    fid = "off-docx-rezipped"
    entries = zip_entries(ctx.docx)
    order = sorted(n for n in entries if n != "[Content_Types].xml") + ["[Content_Types].xml"]
    ctx.write(fid + ".docx", write_zip(entries, order))
    ctx.manifest(
        id=fid, title="off-docx-lorem-a4 re-zipped with [Content_Types].xml as the last entry",
        file=fid + ".docx", issues=[17824, 18030], catalog=["File-type detection"],
        pages=1, expect="converted", kind="vector", marker=False, aspect=0.7071,
        notes="Same entries and bytes as off-docx-lorem-a4, only the zip order differs. libmagic "
              "reports such files as generic OOXML or application/zip; bbb-web's Tika detection "
              "(#18030) must still accept it (#17824).")


# --------------------------------------------------------------------------------------
# 36. XLSX
# --------------------------------------------------------------------------------------
def off_xlsx_one_sheet(ctx):
    fid = "off-xlsx-one-sheet"
    wb = Workbook()
    ws = wb.active
    ws.title = "Fruit"
    for row in TABLE:
        ws.append(row)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "portrait"
    buf = io.BytesIO()
    wb.save(buf)
    ctx.write(fid + ".xlsx", normalize_ooxml(buf.getvalue(), "Fruit table"))
    ctx.manifest(
        id=fid, title="openpyxl sheet: 6x4 table with a bold header row, A4 portrait",
        file=fid + ".xlsx", issues=[], catalog=["Office conversion"],
        pages=1, expect="converted", kind="any", marker=False, aspect=0.7071,
        notes="Header row plus five data rows (strings, integers, decimals). openpyxl writes "
              "[Content_Types].xml last and stamps the current time; the package is rewritten "
              "sorted with fixed dates. LibreOffice 25.8: one A4 page, no <image or <filter.")


# --------------------------------------------------------------------------------------
# 37-40. OpenDocument
# --------------------------------------------------------------------------------------
def _odf_page_layout(doc, name, width, height, orientation, margin="2cm"):
    pl = PageLayout(name=name)
    pl.addElement(PageLayoutProperties(pagewidth=width, pageheight=height, printorientation=orientation,
                                       margin=margin))
    doc.automaticstyles.addElement(pl)
    return pl


def off_odt_one_page(ctx):
    fid = "off-odt-one-page"
    doc = OpenDocumentText()
    pl = _odf_page_layout(doc, "pcfA4", "21cm", "29.7cm", "portrait")
    doc.masterstyles.addElement(MasterPage(name="Standard", pagelayoutname=pl))
    h1 = Style(name="Heading_20_1", displayname="Heading 1", family="paragraph")
    h1.addElement(TextProperties(fontsize="20pt", fontweight="bold"))
    h1.addElement(ParagraphProperties(marginbottom="0.4cm"))
    doc.styles.addElement(h1)
    doc.text.addElement(H(outlinelevel=1, stylename=h1, text="ODT heading"))
    for text in LOREM_PARAS:
        doc.text.addElement(P(text=text))
    ctx.write(fid + ".odt", _odf_bytes(doc))
    ctx.manifest(
        id=fid, title="odfpy text document: a heading and three paragraphs on A4",
        file=fid + ".odt", issues=[], catalog=["Office conversion"],
        pages=1, expect="converted", kind="any", marker=False, aspect=0.7071,
        notes="Minimal ODF text package from odfpy (no settings.xml, no thumbnail). LibreOffice 25.8: one A4 "
              "page, no <image or <filter.")


def _graphic_style(doc, name, rgb):
    s = Style(name=name, family="graphic")
    s.addElement(GraphicProperties(fill="solid", fillcolor=_hex(rgb), stroke="none"))
    doc.automaticstyles.addElement(s)
    return s


def off_odp_one_slide(ctx):
    fid = "off-odp-one-slide"
    doc = OpenDocumentPresentation()
    pl = _odf_page_layout(doc, "PM1", "28cm", "15.75cm", "landscape", margin="0cm")
    mp = MasterPage(name="Default", pagelayoutname=pl)
    doc.masterstyles.addElement(mp)
    page = DrawPage(name="Slide1", masterpagename=mp)
    doc.presentation.addElement(page)
    w, h = 28.0, 15.75
    bx, by, gap = w / 7, h / 7, h / 80
    for name, rgb, (x, y, ww, hh) in (("red", RED, (0, 0, w, by)), ("blue", BLUE, (0, h - by, w, by)),
                                      ("green", GREEN, (0, by + gap, bx, h - 2 * by - 2 * gap)),
                                      ("yellow", YELLOW, (w - bx, by + gap, bx, h - 2 * by - 2 * gap))):
        page.addElement(Rect(stylename=_graphic_style(doc, "gr-" + name, rgb), x="%.3fcm" % x,
                             y="%.3fcm" % y, width="%.3fcm" % ww, height="%.3fcm" % hh))
    tstyle = Style(name="P-title", family="paragraph")
    tstyle.addElement(TextProperties(fontsize="32pt", fontweight="bold"))
    doc.automaticstyles.addElement(tstyle)
    nostroke = Style(name="gr-text", family="graphic")
    nostroke.addElement(GraphicProperties(fill="none", stroke="none"))
    doc.automaticstyles.addElement(nostroke)
    frame = Frame(stylename=nostroke, x="%.3fcm" % (bx + 0.5), y="%.3fcm" % (by + 1.0),
                  width="%.3fcm" % (w - 2 * bx - 1.0), height="5cm")
    box = TextBox()
    box.addElement(P(stylename=tstyle, text="ODP marker slide"))
    box.addElement(P(text="A text frame and four coloured rectangles."))
    frame.addElement(box)
    page.addElement(frame)
    ctx.write(fid + ".odp", _odf_bytes(doc))
    ctx.manifest(
        id=fid, title="odfpy 16:9 presentation slide laid out as the marker",
        file=fid + ".odp", issues=[], catalog=["Office conversion"],
        pages=1, expect="converted", kind="any", marker=True, aspect=1.7778,
        notes="28 x 15.75 cm page; red, blue, green and yellow draw:rect bands without outline, and a "
              "text frame in the middle. LibreOffice 25.8: one 793.7x446.5 pt page with the marker upright, no "
              "<filter.")


def off_ods_one_sheet(ctx):
    fid = "off-ods-one-sheet"
    doc = OpenDocumentSpreadsheet()
    pl = _odf_page_layout(doc, "pm1", "21cm", "29.7cm", "portrait")
    doc.masterstyles.addElement(MasterPage(name="Default", pagelayoutname=pl))
    tstyle = Style(name="ta1", family="table", masterpagename="Default")
    tstyle.addElement(TableProperties(display="true"))
    doc.automaticstyles.addElement(tstyle)
    bold = Style(name="ce1", family="table-cell")
    bold.addElement(TextProperties(fontweight="bold"))
    doc.automaticstyles.addElement(bold)
    table = Table(name="Fruit", stylename=tstyle)
    table.addElement(TableColumn(numbercolumnsrepeated=4))
    for r, row in enumerate(TABLE):
        tr = TableRow()
        for v in row:
            if isinstance(v, str):
                tc = TableCell(valuetype="string", stylename=bold) if r == 0 else TableCell(valuetype="string")
            else:
                tc = TableCell(valuetype="float", value=str(v))
            tc.addElement(P(text=str(v)))
            tr.addElement(tc)
        table.addElement(tr)
    doc.spreadsheet.addElement(table)
    ctx.write(fid + ".ods", _odf_bytes(doc))
    ctx.manifest(
        id=fid, title="odfpy spreadsheet: 6x4 table with a bold header row, A4 portrait",
        file=fid + ".ods", issues=[], catalog=["Office conversion"],
        pages=1, expect="converted", kind="any", marker=False, aspect=0.7071,
        notes="Same table as off-xlsx-one-sheet. LibreOffice 25.8: one A4 page, no <image or <filter.")


def off_odg_one_page(ctx):
    fid = "off-odg-one-page"
    doc = OpenDocumentDrawing()
    pl = _odf_page_layout(doc, "PM1", "29.7cm", "21cm", "landscape", margin="1cm")
    mp = MasterPage(name="Default", pagelayoutname=pl)
    doc.masterstyles.addElement(mp)
    page = DrawPage(name="page1", masterpagename=mp)
    doc.drawing.addElement(page)
    page.addElement(Rect(stylename=_graphic_style(doc, "gr1", (230, 0, 0)), x="2cm", y="2cm",
                         width="8cm", height="5cm"))
    page.addElement(Ellipse(stylename=_graphic_style(doc, "gr2", (0, 0, 230)), x="12cm", y="2cm",
                            width="7cm", height="7cm"))
    line = Style(name="gr3", family="graphic")
    line.addElement(GraphicProperties(stroke="solid", strokecolor="#000000", strokewidth="0.2cm"))
    doc.automaticstyles.addElement(line)
    page.addElement(Line(stylename=line, x1="2cm", y1="12cm", x2="27cm", y2="12cm"))
    nostroke = Style(name="gr4", family="graphic")
    nostroke.addElement(GraphicProperties(fill="none", stroke="none"))
    doc.automaticstyles.addElement(nostroke)
    frame = Frame(stylename=nostroke, x="2cm", y="14cm", width="20cm", height="3cm")
    box = TextBox()
    box.addElement(P(text="ODG drawing: rectangle, circle, line and text"))
    frame.addElement(box)
    page.addElement(frame)
    ctx.write(fid + ".odg", _odf_bytes(doc))
    ctx.manifest(
        id=fid, title="odfpy drawing: a red rectangle, a blue circle, a line and a text frame",
        file=fid + ".odg", issues=[20092], catalog=["File-type detection"],
        pages=1, expect="measure", kind="any", marker=False, aspect=1.4143,
        notes="A4 landscape drawing. .odg was refused as an unsupported type (#20092) and is now in "
              "bbb-web's office list; whether its MIME type passes validation is what this records. "
              "LibreOffice 25.8 converts it to one A4 landscape page.")


# --------------------------------------------------------------------------------------
# 42. Phone photo in a deck
# --------------------------------------------------------------------------------------
def off_pptx_phone_photo_exif6(ctx):
    fid = "off-pptx-phone-photo-exif6"
    prs = Presentation()                              # python-pptx default: 10 x 7.5 in (4:3)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # "Blank"
    pic = slide.shapes.add_picture(io.BytesIO(camera_exif6_jpeg(ctx)), Inches(3), Inches(1),
                                   height=Inches(5))
    pic.left = (prs.slide_width - pic.width) // 2
    pic.top = (prs.slide_height - pic.height) // 2
    buf = io.BytesIO()
    prs.save(buf)
    ctx.write(fid + ".pptx", normalize_ooxml(buf.getvalue(), "Phone photo deck"))
    ctx.manifest(
        id=fid, title="python-pptx 4:3 deck holding the EXIF-6 phone photo",
        file=fid + ".pptx", issues=[26001], catalog=["EXIF orientation, photo in an Office deck"],
        pages=1, expect="converted", kind="any", marker=True, aspect=1.3333,
        notes="Port of deck-camera-exif6.pptx, with the picture centred instead of overflowing the "
              "right edge. python-pptx sizes the frame from the stored 480x320 pixels and ignores "
              "EXIF, and LibreOffice draws the stored pixels, so the slide shows the photo sideways: "
              "known to fail the marker check. Compare with PowerPoint before calling it a bug "
              "(#26001 analysis). LibreOffice 25.8: one 720x540 page, photo turned 90 degrees "
              "counter-clockwise (red band on the left).")


ALL = [off_rtf_one_page, off_pptx_one_slide, off_docx_lorem_a4, off_xlsx_one_sheet, off_odt_one_page,
       off_odp_one_slide, off_ods_one_sheet, off_odg_one_page, off_docx_rezipped,
       off_pptx_phone_photo_exif6]
